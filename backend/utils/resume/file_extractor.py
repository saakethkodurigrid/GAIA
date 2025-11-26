"""
File Extractor for PDF and DOCX resume files.
Extracts text content from resume files.
"""
import io
from typing import Optional
from fastapi import UploadFile


class FileExtractor:
    """Extracts text from PDF and DOCX files."""
    
    def extract_text(self, file: UploadFile) -> Optional[str]:
        """
        Extract text from PDF or DOCX file.
        
        Args:
            file: UploadFile object (PDF or DOCX)
            
        Returns:
            Extracted text content or None if extraction fails
        """
        try:
            # Read file content
            content = file.file.read()
            file.file.seek(0)  # Reset file pointer
            
            # Get file extension
            filename = file.filename or ""
            file_extension = filename.lower().split('.')[-1] if '.' in filename else ""
            
            if file_extension == 'pdf':
                return self._extract_from_pdf(content)
            elif file_extension in ['docx', 'doc']:
                return self._extract_from_docx(content)
            else:
                raise ValueError(f"Unsupported file format: {file_extension}. Only PDF and DOCX are supported.")
                
        except Exception as e:
            print(f"Error extracting text from file {file.filename}: {str(e)}")
            return None
    
    def _extract_from_pdf(self, content: bytes) -> str:
        """
        Extract text from PDF file.
        
        Args:
            content: PDF file content as bytes
            
        Returns:
            Extracted text
        """
        try:
            import pdfplumber
            
            text_parts = []
            with pdfplumber.open(io.BytesIO(content)) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text_parts.append(page_text)
            
            return "\n".join(text_parts)
            
        except ImportError:
            # Fallback to PyPDF2 if pdfplumber not available
            try:
                import PyPDF2
                pdf_reader = PyPDF2.PdfReader(io.BytesIO(content))
                text_parts = []
                for page in pdf_reader.pages:
                    text_parts.append(page.extract_text())
                return "\n".join(text_parts)
            except ImportError:
                raise ImportError("Neither pdfplumber nor PyPDF2 is installed. Please install one: pip install pdfplumber or pip install PyPDF2")
    
    def _extract_from_docx(self, content: bytes) -> str:
        """
        Extract text from DOCX file.
        
        Args:
            content: DOCX file content as bytes
            
        Returns:
            Extracted text
        """
        try:
            from docx import Document
            
            doc = Document(io.BytesIO(content))
            text_parts = []
            
            # Extract text from paragraphs
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text_parts.append(cell.text)
            
            return "\n".join(text_parts)
            
        except ImportError:
            raise ImportError("python-docx is not installed. Please install: pip install python-docx")


# Global instance
file_extractor = FileExtractor()

